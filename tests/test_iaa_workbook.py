import json
import sqlite3

from openpyxl import load_workbook

from scripts.iaa_workbook import COLUMNS, FIRST_ROW, build, import_workbook, parse_row

QUEUE_ROW = {"record_id": "r1", "source_id": "s", "document_id": "d", "document_type": "fare_rule",
             "carrier": "VZ", "chunk_id": "c", "clause_text": "Changes condition: 3 hours before departure",
             "chunk_text": "context", "heuristic": {}}


def _cells(**overrides):
    cells = {key: None for _, key, _ in COLUMNS if key}
    cells.update(overrides)
    return cells


def test_parse_row_label_normalises_values():
    errors = []
    p = parse_row(_cells(decision="Label", policy_topic="reissue_change", action="conditional",
                         time_window="3 hours", before_after_departure="before_departure",
                         fare_difference_required="TRUE", penalty_currency="thb", penalty_amount="1,200",
                         eligibility_conditions="unused tickets only\nadults"), QUEUE_ROW, "row 2", errors)
    assert errors == []
    assert p["label_status"] == "labeled" and p["fare_difference_required"] is True
    assert p["penalty_currency"] == "THB" and p["penalty_amount"] == 1200.0
    assert p["eligibility_conditions"] == ["unused tickets only", "adults"]


def test_parse_row_rejects_missing_required_and_bad_choice():
    errors = []
    parse_row(_cells(decision="label", policy_topic="refunds"), QUEUE_ROW, "row 2", errors)
    assert any("policy_topic 'refunds'" in e for e in errors)
    assert any("action is required" in e for e in errors)


def test_parse_row_skip_keeps_tiers_text_in_notes():
    errors = []
    p = parse_row(_cells(decision="skip", tiers_text="n/a", notes="menu text"), QUEUE_ROW, "row 2", errors)
    assert p["label_status"] == "skip" and "menu text" in p["notes"] and "n/a" in p["notes"]


def test_build_then_import_round_trip(tmp_path):
    queue = tmp_path / "q.jsonl"
    queue.write_text(json.dumps(QUEUE_ROW) + "\n", encoding="utf-8")
    guide = tmp_path / "guide.md"
    guide.write_text("# Guide\nrule", encoding="utf-8")
    xlsx = tmp_path / "wb.xlsx"
    assert build(queue, guide, xlsx) == 0

    wb = load_workbook(xlsx)
    ws = wb["Label"]
    col = {key: i for i, (_, key, _) in enumerate(COLUMNS, 1) if key}
    ws.cell(row=FIRST_ROW, column=col["decision"], value="label")
    ws.cell(row=FIRST_ROW, column=col["policy_topic"], value="reissue_change")
    ws.cell(row=FIRST_ROW, column=col["action"], value="conditional")
    wb.save(xlsx)

    db = tmp_path / "b.db"
    assert import_workbook(xlsx, queue, db, force=False) == 0
    status, payload = sqlite3.connect(db).execute("select label_status, payload_json from manual_labels").fetchone()
    assert status == "labeled" and json.loads(payload)["policy_topic"] == "reissue_change"
    assert import_workbook(xlsx, queue, db, force=False) == 1  # refuses to overwrite
